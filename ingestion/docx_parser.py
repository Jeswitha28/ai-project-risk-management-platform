from pathlib import Path
from docx import Document as WordDocument
from models import Document


def parse_docx(file_path):
    path = Path(file_path)
    word_doc = WordDocument(path)

    parts = []

    for para in word_doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in word_doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    full_text = "\n".join(parts)

    if not full_text.strip():
        raise ValueError(f"No text found in {path.name}. The file may be empty.")

    return Document(
        filename=path.name,
        file_type="docx",
        text=full_text,
    )